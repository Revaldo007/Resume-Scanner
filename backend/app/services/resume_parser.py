"""
Text extraction from resume files (PDF / DOCX).

This module deliberately does ONLY text extraction + basic cleaning.
All information extraction (name, skills, education, etc.) lives in
nlp_processor.py so the pipeline stages stay separate and testable.
"""
import io
import re

import fitz  # PyMuPDF
import docx  # python-docx

ALLOWED_EXTENSIONS = {".pdf", ".docx"}


class UnsupportedFileTypeError(Exception):
    pass


class EmptyResumeError(Exception):
    pass


class CorruptedFileError(Exception):
    pass


def get_extension(filename: str) -> str:
    idx = filename.rfind(".")
    return filename[idx:].lower() if idx != -1 else ""


def extract_text_from_pdf(file_bytes: bytes) -> str:
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
    except Exception as exc:
        raise CorruptedFileError(f"Could not open PDF file: {exc}")

    text_parts = []
    try:
        for page in doc:
            text_parts.append(page.get_text())
    finally:
        doc.close()
    return "\n".join(text_parts)


def extract_text_from_docx(file_bytes: bytes) -> str:
    try:
        document = docx.Document(io.BytesIO(file_bytes))
    except Exception as exc:
        raise CorruptedFileError(f"Could not open DOCX file: {exc}")

    paragraphs = [p.text for p in document.paragraphs]

    # Also pull text out of any tables (resumes often use tables for layout)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    paragraphs.append(cell.text)

    return "\n".join(paragraphs)


def clean_text(raw_text: str) -> str:
    text = raw_text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text(filename: str, file_bytes: bytes) -> str:
    """
    Main entry point: dispatches to the correct extractor based on
    file extension, then cleans the result. Raises on unsupported
    types, corrupted files, or empty output.
    """
    ext = get_extension(filename)
    if ext not in ALLOWED_EXTENSIONS:
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{ext}'. Only PDF and DOCX are allowed."
        )

    if ext == ".pdf":
        raw_text = extract_text_from_pdf(file_bytes)
    else:
        raw_text = extract_text_from_docx(file_bytes)

    cleaned = clean_text(raw_text)
    if not cleaned:
        raise EmptyResumeError(
            "No readable text could be extracted from this resume "
            "(it may be a scanned image or empty file)."
        )
    return cleaned
